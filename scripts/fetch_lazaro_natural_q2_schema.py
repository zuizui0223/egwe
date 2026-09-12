from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urljoin
from urllib.request import Request, urlopen

from openpyxl import load_workbook

DOI = "10.5061/dryad.d51c59zzj"
FILENAME = "Visitation&Seedsetdata_Dryad.xlsx"
DRYAD_ROOT = "https://datadryad.org"
API_ROOT = f"{DRYAD_ROOT}/api/v2"
USER_AGENT = "egwe-natural-q2-schema/1.0"


def _request(url: str, accept: str) -> tuple[str, bytes]:
    req = Request(
        url,
        headers={
            "User-Agent": USER_AGENT,
            "Accept": accept,
            "Referer": "https://datadryad.org/",
            "X-API-Version": "2.1.0",
        },
    )
    with urlopen(req, timeout=180) as response:
        return str(response.headers.get("Content-Type", "")), response.read()


def _json(url: str) -> dict:
    _, payload = _request(url, "application/json")
    return json.loads(payload.decode("utf-8"))


def _absolute(href: str) -> str:
    return urljoin(DRYAD_ROOT, href)


def _href(obj: dict, relation: str) -> str | None:
    value = obj.get("_links", {}).get(relation)
    if isinstance(value, dict) and value.get("href"):
        return _absolute(str(value["href"]))
    return None


def _file_id(item: dict) -> int | None:
    if item.get("id") is not None:
        return int(item["id"])
    self_href = _href(item, "self")
    if self_href:
        match = re.search(r"/files/(\d+)$", self_href)
        if match:
            return int(match.group(1))
    return None


def _resolve_locked_file() -> tuple[str, str, dict]:
    dataset_key = quote(f"doi:{DOI}", safe="")
    dataset_url = f"{API_ROOT}/datasets/{dataset_key}"
    dataset = _json(dataset_url)

    version_url = _href(dataset, "stash:version") or _href(dataset, "version")
    if version_url is None:
        versions = _json(f"{dataset_url}/versions")
        embedded = versions.get("_embedded", {}).get("stash:versions", [])
        if not embedded:
            raise RuntimeError("Dryad dataset exposed no public versions")
        version_url = _href(embedded[0], "self")
        if version_url is None and embedded[0].get("id") is not None:
            version_url = f"{API_ROOT}/versions/{embedded[0]['id']}"
    if version_url is None:
        raise RuntimeError("could not resolve Dryad version")

    version = _json(version_url)
    files_url = _href(version, "stash:files") or _href(version, "files")
    if files_url is None:
        version_id = version.get("id")
        if version_id is None:
            match = re.search(r"/versions/(\d+)$", version_url)
            version_id = int(match.group(1)) if match else None
        if version_id is None:
            raise RuntimeError("could not resolve Dryad version id")
        files_url = f"{API_ROOT}/versions/{version_id}/files"

    page = _json(files_url)
    items = page.get("_embedded", {}).get("stash:files", page.get("files", []))
    item = next(
        (
            x
            for x in items
            if str(x.get("path") or x.get("name") or "") == FILENAME
        ),
        None,
    )
    if item is None:
        raise RuntimeError(f"locked file absent: {FILENAME}")
    file_id = _file_id(item)
    if file_id is None:
        raise RuntimeError("could not resolve locked file id")
    return dataset_url, files_url, {"file_id": file_id, "metadata": item}


def _download_candidates(item: dict, file_id: int) -> list[str]:
    candidates: list[str] = []
    for relation in ("stash:download", "download"):
        href = _href(item, relation)
        if href:
            candidates.append(href)
    candidates.extend(
        [
            f"{DRYAD_ROOT}/stash/downloads/file_stream/{file_id}",
            f"{DRYAD_ROOT}/downloads/file_stream/{file_id}",
        ]
    )
    return list(dict.fromkeys(candidates))


def _digest(payload: bytes, item: dict) -> dict:
    observed = hashlib.sha256(payload).hexdigest()
    expected = item.get("digest")
    digest_type = str(item.get("digestType") or "").lower()
    verified: bool | None = None
    expected_text: str | None = None
    if expected:
        expected_text = str(expected).lower().replace("sha256:", "")
        if "sha-256" in digest_type or "sha256" in digest_type or len(expected_text) == 64:
            verified = observed == expected_text
    return {
        "sha256": observed,
        "metadata_digest": expected_text,
        "metadata_digest_verified": verified,
    }


def _header_value(value: object) -> str:
    if value is None:
        return ""
    # Header-only inspection: stringify the declared first-row labels, never any later cell.
    return str(value).strip()


def _inspect_workbook_headers(payload: bytes) -> list[dict]:
    workbook = load_workbook(
        filename=io.BytesIO(payload),
        read_only=True,
        data_only=False,
    )
    result: list[dict] = []
    for sheet in workbook.worksheets:
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        columns = [_header_value(value) for value in first_row]
        result.append(
            {
                "sheet": sheet.title,
                "max_row_metadata": sheet.max_row,
                "max_column_metadata": sheet.max_column,
                "first_row_labels": columns,
            }
        )
    workbook.close()
    return result


def discover(manifest_path: Path) -> dict:
    dataset_url, files_url, record = _resolve_locked_file()
    item = record["metadata"]
    file_id = int(record["file_id"])
    candidates = _download_candidates(item, file_id)
    errors: list[str] = []
    payload: bytes | None = None
    resolved_url: str | None = None
    content_type: str | None = None

    for url in candidates:
        try:
            candidate_type, candidate_payload = _request(
                url,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*",
            )
        except HTTPError as exc:
            errors.append(f"{url}: HTTP {exc.code}")
            continue
        except (URLError, TimeoutError, OSError) as exc:
            errors.append(f"{url}: {type(exc).__name__}")
            continue

        expected_size = item.get("size")
        if expected_size is not None and int(expected_size) != len(candidate_payload):
            errors.append(
                f"{url}: size mismatch expected={int(expected_size)} observed={len(candidate_payload)}"
            )
            continue
        digest = _digest(candidate_payload, item)
        if digest["metadata_digest_verified"] is False:
            errors.append(f"{url}: digest mismatch")
            continue
        payload = candidate_payload
        resolved_url = url
        content_type = candidate_type
        break

    base = {
        "dataset_doi": DOI,
        "locked_file": FILENAME,
        "dataset_metadata_url": dataset_url,
        "files_metadata_url": files_url,
        "file_id": file_id,
        "metadata_size": item.get("size"),
        "inspection_boundary": (
            "Only Dryad metadata, workbook sheet names/dimensions, and the first/header row of each sheet were inspected. "
            "No cell below the first row, predictor value, seed-set value, effect direction, correlation, model fit, or p-value was read or computed."
        ),
        "preregistration": "manuscript/NEE_Q2_NATURAL_PROSPECTIVE_VALIDATION_PREREGISTRATION.md",
    }

    if payload is None:
        result = {
            **base,
            "status": "raw_byte_access_blocked_before_header_inspection",
            "access_candidates_attempted": candidates,
            "access_errors": errors,
            "next_gate": "Access STOP. Do not substitute another endpoint or analyse published summary statistics as row-level validation.",
        }
    else:
        result = {
            **base,
            "status": "schema_only_discovery_complete",
            "resolved_url": resolved_url,
            "content_type": content_type,
            "bytes": len(payload),
            **_digest(payload, item),
            "sheets": _inspect_workbook_headers(payload),
            "access_candidates_attempted": candidates,
            "next_gate": (
                "Apply the preregistered Stage-A eligibility rule from headers/source definitions only. "
                "If it passes, commit the exact Stage-B model/scoring contract before any response value is read."
            ),
        }

    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--manifest",
        default="artifacts/empirical/lazaro_natural_q2_schema.json",
    )
    args = parser.parse_args()
    result = discover(Path(args.manifest))
    print(
        json.dumps(
            {
                "status": result["status"],
                "locked_file": result["locked_file"],
                "sheets": [
                    {
                        "sheet": sheet["sheet"],
                        "first_row_labels": sheet["first_row_labels"],
                    }
                    for sheet in result.get("sheets", [])
                ],
            },
            sort_keys=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
